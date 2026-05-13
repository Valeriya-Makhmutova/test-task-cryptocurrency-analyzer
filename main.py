import argparse
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from app.endpoints import get_crypto

def main():
    parser = argparse.ArgumentParser(
        description="Консольная утилита для анализа криптовалюты"
    )

    parser.add_argument(
        "-c", "--currency", 
        type=str, 
        default="usd", 
        help="Валюта сравнения (usd, eur, rub)"
    )

    parser.add_argument(
        "--limit", 
        type=int, 
        default=25, 
        help="Количество разных криптовалют в выгрузке (по умолчанию 25)"
    )

    parser.add_argument(
        "--marketCapLine", 
        type=int, 
        default=1000000000, 
        help="Пороговое значение капитализации рынка"
    )

    args = parser.parse_args()
    data = get_crypto(args.currency, args.limit, args.marketCapLine)
    # print('data: ', data)

    new_data = []
    for currency in data:
        short_curr_data = {
            "name": currency.get("name"),
            "symbol": currency.get("symbol"),
            "current_price": currency.get("current_price"),
            "market_cap": currency.get("market_cap")
        }
        new_data.append(short_curr_data)

    # print(new_data)
    names = [currency["name"] for currency in new_data]
    symbols = [currency["symbol"] for currency in new_data]
    prices = [currency["current_price"] for currency in new_data]
    market_caps = [currency["market_cap"] for currency in new_data]

    frame_data = pd.DataFrame({
        "Name": names,
        "Symbol": symbols,
        "Price": prices,
        "Market Capitalization": market_caps,
    })

    frame_data.to_excel('output.xlsx', index=True)
    
    work_book = load_workbook('output.xlsx')
    work_sheet = work_book.active

    green_fill = PatternFill(start_color="008080", end_color="008080", fill_type="solid")
    bold_font = Font(bold=True)

    for cell in work_sheet[1]:
        cell.fill = green_fill
        cell.font = bold_font
    
    for col in work_sheet.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        work_sheet.column_dimensions[col_letter].width = max_len + 3 
    
    with pd.ExcelWriter('output.xlsx', engine='openpyxl') as writer:
        frame_data.to_excel(writer, index=True)

        numbers_format = '#,##0'

        for row in range(2, work_sheet.max_row + 1):
            work_sheet[f'D{row}'].number_format = numbers_format
            work_sheet[f'E{row}'].number_format = numbers_format
    print(frame_data.dtypes)
    work_book.save('output.xlsx')


# try:

# except *название исключения*:
if __name__ == '__main__':
    main()